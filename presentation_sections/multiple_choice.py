"""
Multiple Choice Results Section
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from fpdf import FPDF

from .base import BaseSection


# Hardcoded multiple choice data (same as app.py)
MULTIPLE_CHOICE_DATA = {
    'Future Roles': {
        'question': 'Q3: Which roles do you see presales evolving into?',
        'options': [
            ('Trusted Advisor / Strategic Consultant', 62),
            ('AI-Enhanced Technical Expert', 48),
            ('Customer Success Partner', 45),
            ('Innovation Catalyst', 38),
            ('Digital Transformation Guide', 35),
            ('Value Engineer', 30),
        ]
    },
    'Key Skillsets': {
        'question': 'Q4: What skills will be most important?',
        'options': [
            ('Business Acumen & Strategic Thinking', 58),
            ('AI/ML Literacy', 52),
            ('Consultative Selling', 47),
            ('Data Analytics', 44),
            ('Change Management', 36),
        ]
    }
}


class MultipleChoiceSection(BaseSection):
    """
    Generates Multiple Choice Results slides.

    Shows Q3 (Future Roles) and Q4 (Key Skillsets) vote distributions.
    Uses hardcoded data matching the dashboard.
    """

    def get_slide_count(self) -> int:
        return 2

    def generate_pptx_slides(self, prs: Presentation) -> Presentation:
        # Slide 1: Future Roles
        slide1 = self.add_content_slide(prs, "Future Roles (Q3)")
        self._add_bar_chart(slide1, MULTIPLE_CHOICE_DATA['Future Roles'])

        # Slide 2: Key Skillsets
        slide2 = self.add_content_slide(prs, "Key Skillsets (Q4)")
        self._add_bar_chart(slide2, MULTIPLE_CHOICE_DATA['Key Skillsets'])

        return prs

    def _add_bar_chart(self, slide, data: dict):
        """Add a horizontal bar chart to slide."""
        options = data['options']
        max_value = max(v for _, v in options)

        y_pos = 1.5
        bar_height = 0.6
        max_bar_width = 8.0

        for label, value in options:
            # Label
            self.add_text_box(
                slide, 0.5, y_pos, 4.0, bar_height,
                label, font_size=11, color='text'
            )

            # Bar
            bar_width = (value / max_value) * max_bar_width
            bar = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(4.5), Inches(y_pos + 0.1),
                Inches(bar_width), Inches(bar_height - 0.2)
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = self.colors['secondary']
            bar.line.fill.background()

            # Value
            self.add_text_box(
                slide, 4.5 + bar_width + 0.2, y_pos, 1.0, bar_height,
                str(value), font_size=11, bold=True, color='primary'
            )

            y_pos += bar_height + 0.15

    def generate_pdf_content(self, pdf: FPDF) -> FPDF:
        # Page 1: Future Roles
        self.pdf_add_content_page(pdf, "Future Roles (Q3)")
        self._add_pdf_bars(pdf, MULTIPLE_CHOICE_DATA['Future Roles'])

        # Page 2: Key Skillsets
        self.pdf_add_content_page(pdf, "Key Skillsets (Q4)")
        self._add_pdf_bars(pdf, MULTIPLE_CHOICE_DATA['Key Skillsets'])

        return pdf

    def _add_pdf_bars(self, pdf: FPDF, data: dict):
        """Add horizontal bars to PDF."""
        options = data['options']
        max_value = max(v for _, v in options)

        pdf.ln(5)

        for label, value in options:
            # Label
            pdf.set_font('Helvetica', '', 10)
            pdf.set_text_color(*self.pdf_colors['text'])
            pdf.cell(80, 8, label)

            # Bar
            bar_width = (value / max_value) * 120
            pdf.set_fill_color(*self.pdf_colors['secondary'])
            pdf.cell(bar_width, 8, '', fill=True)

            # Value
            pdf.set_font('Helvetica', 'B', 10)
            pdf.set_text_color(*self.pdf_colors['primary'])
            pdf.cell(0, 8, f"  {value}", new_x='LMARGIN', new_y='NEXT')

        pdf.set_text_color(*self.pdf_colors['text'])
