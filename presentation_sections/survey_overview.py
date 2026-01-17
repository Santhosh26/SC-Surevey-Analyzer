"""
Survey Overview Section - Enhanced with Charts
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from fpdf import FPDF

from .base import BaseSection
from .charts import ChartGenerator


class SurveyOverviewSection(BaseSection):
    """
    Generates the Survey Overview slide showing response statistics.

    Enhanced with:
    - Horizontal bar chart for response distribution
    - Metric cards for key statistics

    Requires: raw-data.csv
    """

    def get_slide_count(self) -> int:
        return 1

    def _get_stats(self) -> dict:
        """Calculate survey statistics."""
        raw_data = self.data.get('raw_data')
        if raw_data is None:
            return {
                'total_responses': 0,
                'total_questions': 0,
                'questions': []
            }

        questions = raw_data['Question'].unique()
        question_counts = raw_data.groupby('Question').size().to_dict()

        # Sort questions by response count
        sorted_questions = sorted(
            [(q, question_counts.get(q, 0)) for q in questions],
            key=lambda x: x[1],
            reverse=True
        )

        return {
            'total_responses': len(raw_data),
            'total_questions': len(questions),
            'questions': [
                {'question': q[:45] + '...' if len(q) > 45 else q, 'count': count}
                for q, count in sorted_questions
            ][:10]  # Top 10 questions
        }

    def generate_pptx_slides(self, prs: Presentation) -> Presentation:
        stats = self._get_stats()
        charts = ChartGenerator()

        slide = self.add_content_slide(prs, "Survey Overview")

        # Generate metric cards
        metrics = [
            (f"{stats['total_responses']:,}", "Total Responses", "primary"),
            (str(stats['total_questions']), "Questions", "secondary"),
            (f"{stats['total_responses'] // max(stats['total_questions'], 1)}", "Avg/Question", "accent"),
        ]
        metrics_path = charts.metric_cards(metrics, figsize=(9, 1.8))
        self.add_chart_image(slide, metrics_path, 0.3, 1.2, 9, 1.5)

        # Generate response distribution bar chart
        if stats['questions']:
            chart_data = [(q['question'], q['count']) for q in stats['questions']]
            chart_path = charts.horizontal_bar_chart(
                chart_data,
                title="Responses per Question",
                color='primary',
                max_items=8,
                figsize=(8, 4.2)
            )
            self.add_chart_image(slide, chart_path, 0.3, 2.9, 8.5, 4.3)

        # Data source note
        self.add_text_box(
            slide, 9.0, 6.5, 4.0, 0.5,
            "Source: Mentimeter Survey",
            font_size=9, color='neutral', align='right'
        )

        charts.cleanup()
        return prs

    def generate_pdf_content(self, pdf: FPDF) -> FPDF:
        stats = self._get_stats()
        charts = ChartGenerator()

        self.pdf_add_content_page(pdf, "Survey Overview")

        # Key metrics
        pdf.set_font('Helvetica', 'B', 24)
        pdf.set_text_color(*self.pdf_colors['primary'])
        pdf.cell(90, 15, f"{stats['total_responses']:,}", align='C')
        pdf.set_text_color(*self.pdf_colors['secondary'])
        pdf.cell(90, 15, str(stats['total_questions']), align='C', new_x='LMARGIN', new_y='NEXT')

        pdf.set_font('Helvetica', '', 10)
        pdf.set_text_color(*self.pdf_colors['neutral'])
        pdf.cell(90, 6, "Total Responses", align='C')
        pdf.cell(90, 6, "Questions Analyzed", align='C', new_x='LMARGIN', new_y='NEXT')

        pdf.ln(8)

        # Response distribution (simple bar representation)
        pdf.set_font('Helvetica', 'B', 12)
        pdf.set_text_color(*self.pdf_colors['primary'])
        pdf.cell(0, 8, "Responses per Question:", new_x='LMARGIN', new_y='NEXT')
        pdf.ln(2)

        max_count = max([q['count'] for q in stats['questions']]) if stats['questions'] else 1

        for item in stats['questions'][:8]:
            # Question label
            pdf.set_font('Helvetica', '', 9)
            pdf.set_text_color(*self.pdf_colors['text'])
            label = item['question'][:40] + '...' if len(item['question']) > 40 else item['question']
            pdf.cell(100, 6, label)

            # Bar
            bar_width = (item['count'] / max_count) * 80
            pdf.set_fill_color(*self.pdf_colors['secondary'])
            pdf.cell(bar_width, 6, '', fill=True)

            # Count
            pdf.set_font('Helvetica', 'B', 9)
            pdf.cell(0, 6, f"  {item['count']}", new_x='LMARGIN', new_y='NEXT')

        charts.cleanup()
        return pdf
