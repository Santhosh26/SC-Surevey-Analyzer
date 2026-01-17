"""
Action Plan Section - Enhanced with Timeline
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from fpdf import FPDF

from .base import BaseSection
from .charts import ChartGenerator


class ActionPlanSection(BaseSection):
    """
    Generates the Recommended Action Plan slide from LLM summaries.

    Enhanced with:
    - Visual timeline showing action phases
    - Clear phase separation

    Requires: llm_summaries.json (overall_summary.action_plan)

    Shows actions organized by timeframe:
    - 0-30 days (Quick wins)
    - 30-90 days (Medium term)
    - 90+ days (Long term)
    """

    def get_slide_count(self) -> int:
        return 1

    def _get_action_plan(self) -> dict:
        """Extract action plan from LLM data."""
        llm_data = self.data.get('llm_summaries', {})
        overall = llm_data.get('overall_summary', {})
        actions = overall.get('action_plan', [])

        # Group by timeframe
        phases = {
            '0-30 days': [],
            '30-90 days': [],
            '90+ days': [],
        }

        if not actions:
            phases['0-30 days'] = ['Define quick wins']
            phases['30-90 days'] = ['Plan medium-term goals']
            phases['90+ days'] = ['Execute long-term strategy']
        else:
            for action in actions:
                tf = action.get('timeframe', '90+ days')
                action_text = action.get('action', 'Action item')
                if tf in phases:
                    phases[tf].append(action_text)

        return phases

    def generate_pptx_slides(self, prs: Presentation) -> Presentation:
        phases = self._get_action_plan()
        charts = ChartGenerator()

        slide = self.add_content_slide(prs, "Recommended Action Plan")

        # Generate timeline visualization
        timeline_path = charts.timeline(phases, figsize=(12, 4.5))
        self.add_chart_image(slide, timeline_path, 0.3, 1.3, 12.5, 4.5)

        # Summary footer
        total_actions = sum(len(v) for v in phases.values())
        quick_wins = len(phases.get('0-30 days', []))

        self.add_text_box(
            slide, 0.5, 6.2, 6.0, 0.4,
            f"Total Actions: {total_actions}",
            font_size=11, bold=True, color='primary'
        )

        self.add_text_box(
            slide, 6.5, 6.2, 6.0, 0.4,
            f"Quick Wins (0-30 days): {quick_wins}",
            font_size=11, bold=True, color='positive'
        )

        # Key insight
        self.add_text_box(
            slide, 0.5, 6.7, 12.0, 0.5,
            "Focus on quick wins first to build momentum, then scale systematically.",
            font_size=10, color='neutral', align='center'
        )

        charts.cleanup()
        return prs

    def generate_pdf_content(self, pdf: FPDF) -> FPDF:
        llm_data = self.data.get('llm_summaries', {})
        overall = llm_data.get('overall_summary', {})
        actions = overall.get('action_plan', [])

        self.pdf_add_content_page(pdf, "Recommended Action Plan")

        # Group actions by timeframe
        timeframes = {
            '0-30 days': {'color': self.pdf_colors['positive'], 'label': 'NOW', 'actions': []},
            '30-90 days': {'color': self.pdf_colors['accent'], 'label': 'NEXT', 'actions': []},
            '90+ days': {'color': self.pdf_colors['secondary'], 'label': 'LATER', 'actions': []},
        }

        for action in actions:
            tf = action.get('timeframe', '90+ days')
            if tf in timeframes:
                timeframes[tf]['actions'].append(action)

        for tf_name, tf_data in timeframes.items():
            # Timeframe header with visual indicator
            pdf.set_font('Helvetica', 'B', 12)
            pdf.set_fill_color(*tf_data['color'])
            pdf.set_text_color(*self.pdf_colors['white'])
            pdf.cell(25, 8, f"  {tf_data['label']}", fill=True)
            pdf.set_fill_color(*self.pdf_colors['background'])
            pdf.set_text_color(*tf_data['color'])
            pdf.cell(0, 8, f"  {tf_name}", fill=True, new_x='LMARGIN', new_y='NEXT')

            pdf.set_text_color(*self.pdf_colors['text'])

            for action in tf_data['actions']:
                action_text = action.get('action', '')

                # Action item with bullet - use proper indentation
                pdf.set_font('Helvetica', '', 10)
                pdf.set_x(pdf.l_margin + 5)
                pdf.multi_cell(0, 6, f"- {self._sanitize_text(action_text)}", new_x='LMARGIN', new_y='NEXT')

            if not tf_data['actions']:
                pdf.set_font('Helvetica', 'I', 9)
                pdf.set_text_color(*self.pdf_colors['neutral'])
                pdf.cell(0, 6, "  No actions defined", new_x='LMARGIN', new_y='NEXT')
                pdf.set_text_color(*self.pdf_colors['text'])

            pdf.ln(4)

        # Summary
        total_actions = sum(len(tf['actions']) for tf in timeframes.values())
        pdf.ln(5)
        pdf.set_font('Helvetica', 'B', 10)
        pdf.set_text_color(*self.pdf_colors['primary'])
        pdf.cell(0, 8, f"Total Actions: {total_actions}", align='C')

        return pdf
